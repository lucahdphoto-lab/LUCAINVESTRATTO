"""
Institutional Excel report generator for LUCAINVESTRATTO.
Creates a professional multi-sheet formatted workbook (.xlsx) using OpenPyXL.
Ensures ALL numeric data is strictly typed as numeric (float/int), NEVER as formatted text.
Standard Excel number formats, active formulas, and explicit RIGHT ALIGNMENT are enforced
on all numerical cells, guaranteeing seamless mathematical operations and institutional presentation.
"""
import io
from typing import List, Dict, Any, Optional
from datetime import datetime
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from lucainvestratto.models import BankAccount, Holding, Transaction
from lucainvestratto.analytics.cashflow import calculate_cashflow_summary
from lucainvestratto.analytics.assets import calculate_asset_summary
from lucainvestratto.analytics.risk import calculate_risk_metrics

# Module-level Typography & Colors
FONT_TITLE = Font(name="Calibri", size=16, bold=True, color="0F172A")
FONT_SUBTITLE = Font(name="Calibri", size=10, italic=True, color="64748B")
FONT_HEADER = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
FONT_BOLD = Font(name="Calibri", size=11, bold=True, color="1E293B")
FONT_NORMAL = Font(name="Calibri", size=10, color="1E293B")
FONT_TOTAL = Font(name="Calibri", size=11, bold=True, color="0F172A")

FILL_HEADER = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
FILL_ALT = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
FILL_TOTAL = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")

THIN_SIDE = Side(style='thin', color='CBD5E1')
DOUBLE_BOTTOM_SIDE = Side(style='double', color='0F172A')
THICK_TOP_SIDE = Side(style='medium', color='0F172A')

THIN_BORDER = Border(left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=THIN_SIDE)
TOTAL_BORDER = Border(left=THIN_SIDE, right=THIN_SIDE, top=THICK_TOP_SIDE, bottom=DOUBLE_BOTTOM_SIDE)

# Standard Excel Number Formats
FMT_CURRENCY = '#,##0.00 "€";[Red]-#,##0.00 "€";0.00 "€"'
FMT_CURRENCY_PLAIN = '#,##0.00 "€"'
FMT_PCT_SIGNED = '+0.00%;-0.00%;0.00%'
FMT_PCT = '0.00%'
FMT_QTY = '#,##0.0000'
FMT_DECIMAL = '0.00'

# Explicit Alignments: Numbers MUST always be right-aligned!
ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")
ALIGN_HEADER_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
ALIGN_HEADER_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
ALIGN_HEADER_RIGHT = Alignment(horizontal="right", vertical="center", wrap_text=True)


def _to_num(val: Any, default: float = 0.0) -> float:
    """Ensures value is a native float without NaN, None, or string issues."""
    if val is None:
        return default
    try:
        f = float(val)
        import math
        if math.isnan(f) or math.isinf(f):
            return default
        return f
    except (ValueError, TypeError):
        return default


def generate_excel_report(
    accounts: List[BankAccount],
    holdings: List[Holding],
    transactions: List[Transaction],
    account_filter_name: str = "Tutti i conti (Consolidato)",
    benchmark_code: str = "MSCI_WORLD"
) -> bytes:
    wb = openpyxl.Workbook()

    # Local aliases for shared constants
    font_title = FONT_TITLE
    font_subtitle = FONT_SUBTITLE
    font_header = FONT_HEADER
    font_bold = FONT_BOLD
    font_normal = FONT_NORMAL
    font_total = FONT_TOTAL

    fill_header = FILL_HEADER
    fill_alt = FILL_ALT
    fill_total = FILL_TOTAL

    thin_border = THIN_BORDER
    total_border = TOTAL_BORDER

    align_left = ALIGN_LEFT
    align_center = ALIGN_CENTER
    align_right = ALIGN_RIGHT
    align_header_left = ALIGN_HEADER_LEFT
    align_header_center = ALIGN_HEADER_CENTER
    align_header_right = ALIGN_HEADER_RIGHT

    # Compute analytics
    cash_sum = calculate_cashflow_summary(transactions, accounts)
    asset_sum = calculate_asset_summary(holdings)
    risk_sum = calculate_risk_metrics(holdings, benchmark_code=benchmark_code)

    tot_liquidita = _to_num(cash_sum.get("final_balance", 0.0))
    tot_investito = _to_num(asset_sum.get("total_market_value", 0.0))
    patrimonio_totale = tot_liquidita + tot_investito

    acc_map = {a.id: a.name for a in accounts}

    # ================= SHEET 1: SINTESI & METRICHE =================
    ws1 = wb.active
    ws1.title = "Sintesi & Metriche Rischio"
    ws1.views.sheetView[0].showGridLines = True

    ws1["A1"] = "LUCAINVESTRATTO - REPORT ESECUTIVO & ANALISI PROFESSIONALE"
    ws1["A1"].font = font_title
    ws1["A2"] = f"Generato il: {datetime.now().strftime('%d/%m/%Y %H:%M')} | Perimetro: {account_filter_name} | Benchmark: {benchmark_code}"
    ws1["A2"].font = font_subtitle

    # KPI Table Headers
    ws1["A4"] = "INDICATORE PATRIMONIALE"
    ws1["A4"].alignment = align_header_left
    ws1["B4"] = "VALORE (€)"
    ws1["B4"].alignment = align_header_right
    ws1["D4"] = "METRICA DI RISCHIO FACTSET"
    ws1["D4"].alignment = align_header_left
    ws1["E4"] = "VALORE"
    ws1["E4"].alignment = align_header_right

    for col in ["A4", "B4", "D4", "E4"]:
        ws1[col].font = font_header
        ws1[col].fill = fill_header

    kpi_left = [
        ("Patrimonio Netto Totale", patrimonio_totale, FMT_CURRENCY),
        ("Liquidità Totale Conti", tot_liquidita, FMT_CURRENCY),
        ("Controvalore Portafoglio", tot_investito, FMT_CURRENCY),
        ("Costo Storico (PMC)", _to_num(asset_sum.get("total_invested", 0.0)), FMT_CURRENCY),
        ("Plus/Minusvalenza Latente", _to_num(asset_sum.get("total_unrealized_pl", 0.0)), FMT_CURRENCY),
        ("Plus/Minusvalenza %", _to_num(asset_sum.get("total_unrealized_pl_pct", 0.0)) / 100.0, FMT_PCT_SIGNED),
        ("Flusso Netto Cassa Periodo", _to_num(cash_sum.get("net_cashflow", 0.0)), FMT_CURRENCY),
        ("Tasso di Risparmio", _to_num(cash_sum.get("savings_rate_pct", 0.0)) / 100.0, FMT_PCT),
        ("Dividendi Annui Stimati", _to_num(asset_sum.get("expected_annual_income", 0.0)), FMT_CURRENCY),
        ("Dividend Yield di Portafoglio", _to_num(asset_sum.get("portfolio_dividend_yield", 0.0)) / 100.0, FMT_PCT)
    ]

    kpi_right = [
        ("Indice di Sharpe (Ann.)", _to_num(risk_sum.get("sharpe_ratio", 0.0)), FMT_DECIMAL),
        ("Indice di Sortino", _to_num(risk_sum.get("sortino_ratio", 0.0)), FMT_DECIMAL),
        ("Indice di Calmar", _to_num(risk_sum.get("calmar_ratio", 0.0)), FMT_DECIMAL),
        ("Indice di Treynor", _to_num(risk_sum.get("treynor_ratio", 0.0)), FMT_DECIMAL),
        ("Rendimento Annualizzato (CAGR)", _to_num(risk_sum.get("cagr", 0.0)) / 100.0, FMT_PCT_SIGNED),
        ("Volatilità Annualizzata", _to_num(risk_sum.get("volatility_annualized", 0.0)) / 100.0, FMT_PCT),
        ("Massimo Drawdown (MDD)", _to_num(risk_sum.get("max_drawdown", 0.0)) / 100.0, FMT_PCT),
        ("Value at Risk (VaR 95% Ann.)", _to_num(risk_sum.get("var_95_annual", 0.0)) / 100.0, FMT_PCT),
        ("Value at Risk (VaR 99% Ann.)", _to_num(risk_sum.get("var_99_annual", 0.0)) / 100.0, FMT_PCT),
        ("Conditional VaR (CVaR 95%)", _to_num(risk_sum.get("cvar_95_annual", 0.0)) / 100.0, FMT_PCT),
        ("Beta vs Benchmark", _to_num(risk_sum.get("beta", 1.0)), FMT_DECIMAL),
        ("Alfa di Jensen", _to_num(risk_sum.get("alpha_jensen", 0.0)) / 100.0, FMT_PCT_SIGNED),
        ("Coefficiente Determinazione (R²)", _to_num(risk_sum.get("r_squared", 0.0)), FMT_DECIMAL),
        ("Tracking Error vs Benchmark", _to_num(risk_sum.get("tracking_error", 0.0)) / 100.0, FMT_PCT),
        ("Information Ratio", _to_num(risk_sum.get("information_ratio", 0.0)), FMT_DECIMAL),
        ("Win Rate Giornaliero", _to_num(risk_sum.get("win_rate", 0.0)) / 100.0, FMT_PCT)
    ]

    for idx, (label, val, fmt) in enumerate(kpi_left, start=5):
        ws1[f"A{idx}"] = label
        ws1[f"A{idx}"].font = font_bold if idx == 5 else font_normal
        ws1[f"A{idx}"].alignment = align_left
        ws1[f"A{idx}"].border = thin_border
        
        c = ws1[f"B{idx}"]
        c.value = _to_num(val)
        c.font = font_bold if idx == 5 else font_normal
        c.number_format = fmt
        c.alignment = align_right  # Explicit Right Alignment for numbers
        c.border = thin_border

    for idx, (label, val, fmt) in enumerate(kpi_right, start=5):
        ws1[f"D{idx}"] = label
        ws1[f"D{idx}"].font = font_normal
        ws1[f"D{idx}"].alignment = align_left
        ws1[f"D{idx}"].border = thin_border
        
        c = ws1[f"E{idx}"]
        c.value = _to_num(val)
        c.font = font_bold if idx in [5, 6, 7] else font_normal
        c.number_format = fmt
        c.alignment = align_right  # Explicit Right Alignment for numbers
        c.border = thin_border

    # ================= SHEET 2: ASSET ALLOCATION =================
    ws2 = wb.create_sheet(title="Asset Allocation")
    ws2.views.sheetView[0].showGridLines = True
    ws2["A1"] = "RIPARTIZIONE PORTAFOGLIO PER ASSET CLASS"
    ws2["A1"].font = font_title

    headers_alloc = [
        ("Asset Class", align_header_left),
        ("Controvalore (€)", align_header_right),
        ("Costo Storico (€)", align_header_right),
        ("Peso (%)", align_header_right),
        ("P&L (€)", align_header_right),
        ("Rendimento (%)", align_header_right)
    ]
    for col_idx, (h, aln) in enumerate(headers_alloc, 1):
        cell = ws2.cell(row=3, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = aln

    df_class = asset_sum.get("df_by_class")
    start_row_alloc = 4
    current_row_alloc = start_row_alloc

    if df_class is not None and not df_class.empty:
        for r in df_class.itertuples():
            row_idx = current_row_alloc
            c1 = ws2.cell(row=row_idx, column=1, value=str(getattr(r, "asset_class", "")))
            c1.font = font_bold
            c1.alignment = align_left
            
            c2 = ws2.cell(row=row_idx, column=2, value=_to_num(getattr(r, "market_value", 0.0)))
            c2.number_format = FMT_CURRENCY_PLAIN
            c2.font = font_normal
            c2.alignment = align_right

            c3 = ws2.cell(row=row_idx, column=3, value=_to_num(getattr(r, "total_cost", 0.0)))
            c3.number_format = FMT_CURRENCY_PLAIN
            c3.font = font_normal
            c3.alignment = align_right

            c4 = ws2.cell(row=row_idx, column=4, value=_to_num(getattr(r, "weight_pct", 0.0)) / 100.0)
            c4.number_format = FMT_PCT
            c4.font = font_normal
            c4.alignment = align_right

            c5 = ws2.cell(row=row_idx, column=5, value=_to_num(getattr(r, "unrealized_pl", 0.0)))
            c5.number_format = FMT_CURRENCY
            c5.font = font_normal
            c5.alignment = align_right

            c6 = ws2.cell(row=row_idx, column=6, value=_to_num(getattr(r, "pl_pct", 0.0)) / 100.0)
            c6.number_format = FMT_PCT_SIGNED
            c6.font = font_normal
            c6.alignment = align_right

            for c in range(1, 7):
                ws2.cell(row=row_idx, column=c).border = thin_border
                if row_idx % 2 == 0:
                    ws2.cell(row=row_idx, column=c).fill = fill_alt
            current_row_alloc += 1

        # Totals Row with active formulas
        tot_row = current_row_alloc
        c_tot_label = ws2.cell(row=tot_row, column=1, value="TOTALE PORTAFOGLIO")
        c_tot_label.font = font_total
        c_tot_label.alignment = align_left
        
        tc2 = ws2.cell(row=tot_row, column=2, value=f"=SUM(B{start_row_alloc}:B{tot_row - 1})")
        tc2.number_format = FMT_CURRENCY_PLAIN
        tc2.font = font_total
        tc2.alignment = align_right

        tc3 = ws2.cell(row=tot_row, column=3, value=f"=SUM(C{start_row_alloc}:C{tot_row - 1})")
        tc3.number_format = FMT_CURRENCY_PLAIN
        tc3.font = font_total
        tc3.alignment = align_right

        tc4 = ws2.cell(row=tot_row, column=4, value=f"=SUM(D{start_row_alloc}:D{tot_row - 1})")
        tc4.number_format = FMT_PCT
        tc4.font = font_total
        tc4.alignment = align_right

        tc5 = ws2.cell(row=tot_row, column=5, value=f"=SUM(E{start_row_alloc}:E{tot_row - 1})")
        tc5.number_format = FMT_CURRENCY
        tc5.font = font_total
        tc5.alignment = align_right

        tc6 = ws2.cell(row=tot_row, column=6, value=f"=IF(C{tot_row}>0, (B{tot_row}-C{tot_row})/C{tot_row}, 0)")
        tc6.number_format = FMT_PCT_SIGNED
        tc6.font = font_total
        tc6.alignment = align_right

        for c in range(1, 7):
            ws2.cell(row=tot_row, column=c).border = total_border
            ws2.cell(row=tot_row, column=c).fill = fill_total

    # ================= SHEET 3: POSIZIONI DETTAGLIATE =================
    ws3 = wb.create_sheet(title="Posizioni Portafoglio")
    ws3.views.sheetView[0].showGridLines = True
    ws3["A1"] = "DETTAGLIO TITOLI ED ASSET IN PORTAFOGLIO"
    ws3["A1"].font = font_title

    headers_pos = [
        ("Simbolo / ISIN", align_header_center),
        ("Nome Strumento", align_header_left),
        ("Asset Class", align_header_left),
        ("Settore", align_header_left),
        ("Geografia", align_header_center),
        ("Valuta", align_header_center),
        ("Quantità", align_header_right),
        ("PMC (€)", align_header_right),
        ("Prezzo Attuale (€)", align_header_right),
        ("Costo Totale (€)", align_header_right),
        ("Controvalore (€)", align_header_right),
        ("P&L (€)", align_header_right),
        ("Rendimento (%)", align_header_right),
        ("Peso (%)", align_header_right),
        ("Yield Dividendo (%)", align_header_right),
        ("Reddito Annuo Stimato (€)", align_header_right)
    ]
    for col_idx, (h, aln) in enumerate(headers_pos, 1):
        cell = ws3.cell(row=3, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = aln

    df_holdings = asset_sum.get("df_holdings")
    start_row_h = 4
    current_row_h = start_row_h

    if df_holdings is not None and not df_holdings.empty:
        for h in df_holdings.itertuples():
            row_idx = current_row_h
            
            c_sym = ws3.cell(row=row_idx, column=1, value=str(getattr(h, "symbol", "")))
            c_sym.font = font_bold
            c_sym.alignment = align_center

            c_name = ws3.cell(row=row_idx, column=2, value=str(getattr(h, "name", "")))
            c_name.font = font_normal
            c_name.alignment = align_left

            c_cls = ws3.cell(row=row_idx, column=3, value=str(getattr(h, "asset_class", "")))
            c_cls.font = font_normal
            c_cls.alignment = align_left

            c_sec = ws3.cell(row=row_idx, column=4, value=str(getattr(h, "sector", "")))
            c_sec.font = font_normal
            c_sec.alignment = align_left

            c_geo = ws3.cell(row=row_idx, column=5, value=str(getattr(h, "geography", "")))
            c_geo.font = font_normal
            c_geo.alignment = align_center

            c_cur = ws3.cell(row=row_idx, column=6, value=str(getattr(h, "currency", "EUR")))
            c_cur.font = font_normal
            c_cur.alignment = align_center

            # Numeric columns with strictly RIGHT alignment
            c_qty = ws3.cell(row=row_idx, column=7, value=_to_num(getattr(h, "quantity", 0.0)))
            c_qty.number_format = FMT_QTY
            c_qty.font = font_normal
            c_qty.alignment = align_right

            c_pmc = ws3.cell(row=row_idx, column=8, value=_to_num(getattr(h, "average_buy_price", 0.0)))
            c_pmc.number_format = FMT_CURRENCY_PLAIN
            c_pmc.font = font_normal
            c_pmc.alignment = align_right

            c_px = ws3.cell(row=row_idx, column=9, value=_to_num(getattr(h, "current_price", 0.0)))
            c_px.number_format = FMT_CURRENCY_PLAIN
            c_px.font = font_normal
            c_px.alignment = align_right

            c_cost = ws3.cell(row=row_idx, column=10, value=_to_num(getattr(h, "total_cost", 0.0)))
            c_cost.number_format = FMT_CURRENCY_PLAIN
            c_cost.font = font_normal
            c_cost.alignment = align_right

            c_mv = ws3.cell(row=row_idx, column=11, value=_to_num(getattr(h, "market_value", 0.0)))
            c_mv.number_format = FMT_CURRENCY_PLAIN
            c_mv.font = font_normal
            c_mv.alignment = align_right

            c_pl = ws3.cell(row=row_idx, column=12, value=_to_num(getattr(h, "unrealized_pl", 0.0)))
            c_pl.number_format = FMT_CURRENCY
            c_pl.font = font_normal
            c_pl.alignment = align_right

            c_pl_pct = ws3.cell(row=row_idx, column=13, value=_to_num(getattr(h, "unrealized_pl_pct", 0.0)) / 100.0)
            c_pl_pct.number_format = FMT_PCT_SIGNED
            c_pl_pct.font = font_normal
            c_pl_pct.alignment = align_right

            c_w = ws3.cell(row=row_idx, column=14, value=_to_num(getattr(h, "weight_pct", 0.0)) / 100.0)
            c_w.number_format = FMT_PCT
            c_w.font = font_normal
            c_w.alignment = align_right

            c_yd = ws3.cell(row=row_idx, column=15, value=_to_num(getattr(h, "annual_dividend_yield", 0.0)) / 100.0)
            c_yd.number_format = FMT_PCT
            c_yd.font = font_normal
            c_yd.alignment = align_right

            c_inc = ws3.cell(row=row_idx, column=16, value=_to_num(getattr(h, "expected_annual_income", 0.0)))
            c_inc.number_format = FMT_CURRENCY_PLAIN
            c_inc.font = font_normal
            c_inc.alignment = align_right

            for c in range(1, 17):
                ws3.cell(row=row_idx, column=c).border = thin_border
                if row_idx % 2 == 0:
                    ws3.cell(row=row_idx, column=c).fill = fill_alt
            current_row_h += 1

        # Totals Row with active formulas
        tot_h_row = current_row_h
        c_tot_h_lbl = ws3.cell(row=tot_h_row, column=1, value="TOTALE GENERALE")
        c_tot_h_lbl.font = font_total
        c_tot_h_lbl.alignment = align_left

        for c_empty in range(2, 10):
            ws3.cell(row=tot_h_row, column=c_empty, value="")

        t_cost = ws3.cell(row=tot_h_row, column=10, value=f"=SUM(J{start_row_h}:J{tot_h_row - 1})")
        t_cost.number_format = FMT_CURRENCY_PLAIN
        t_cost.font = font_total
        t_cost.alignment = align_right

        t_mv = ws3.cell(row=tot_h_row, column=11, value=f"=SUM(K{start_row_h}:K{tot_h_row - 1})")
        t_mv.number_format = FMT_CURRENCY_PLAIN
        t_mv.font = font_total
        t_mv.alignment = align_right

        t_pl = ws3.cell(row=tot_h_row, column=12, value=f"=SUM(L{start_row_h}:L{tot_h_row - 1})")
        t_pl.number_format = FMT_CURRENCY
        t_pl.font = font_total
        t_pl.alignment = align_right

        t_pl_pct = ws3.cell(row=tot_h_row, column=13, value=f"=IF(J{tot_h_row}>0, L{tot_h_row}/J{tot_h_row}, 0)")
        t_pl_pct.number_format = FMT_PCT_SIGNED
        t_pl_pct.font = font_total
        t_pl_pct.alignment = align_right

        t_w = ws3.cell(row=tot_h_row, column=14, value=f"=SUM(N{start_row_h}:N{tot_h_row - 1})")
        t_w.number_format = FMT_PCT
        t_w.font = font_total
        t_w.alignment = align_right

        ws3.cell(row=tot_h_row, column=15, value="")

        t_inc = ws3.cell(row=tot_h_row, column=16, value=f"=SUM(P{start_row_h}:P{tot_h_row - 1})")
        t_inc.number_format = FMT_CURRENCY_PLAIN
        t_inc.font = font_total
        t_inc.alignment = align_right

        for c in range(1, 17):
            ws3.cell(row=tot_h_row, column=c).border = total_border
            ws3.cell(row=tot_h_row, column=c).fill = fill_total

    # ================= SHEET 4: ESTRATTO CONTO =================
    ws4 = wb.create_sheet(title="Estratto Conto Transazioni")
    ws4.views.sheetView[0].showGridLines = True
    ws4["A1"] = "ESTRATTO CONTO E MOVIMENTI BANCARI"
    ws4["A1"].font = font_title

    headers_tx = [
        ("Data", align_header_center),
        ("ID Transazione", align_header_center),
        ("ID Conto", align_header_center),
        ("Nome Conto", align_header_left),
        ("Tipo Operazione", align_header_left),
        ("Categoria", align_header_left),
        ("Importo (€)", align_header_right),
        ("Commissioni (€)", align_header_right),
        ("Valuta", align_header_center),
        ("Descrizione / Causale", align_header_left),
        ("Titolo / Ticker", align_header_center),
        ("Quantità", align_header_right),
        ("Prezzo Unitario (€)", align_header_right)
    ]
    for col_idx, (h, aln) in enumerate(headers_tx, 1):
        cell = ws4.cell(row=3, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = aln

    start_row_tx = 4
    current_row_tx = start_row_tx

    for tx in transactions:
        row_idx = current_row_tx
        
        c_dt = ws4.cell(row=row_idx, column=1, value=str(tx.date or ""))
        c_dt.font = font_normal
        c_dt.alignment = align_center

        c_id = ws4.cell(row=row_idx, column=2, value=str(tx.id or ""))
        c_id.font = font_normal
        c_id.alignment = align_center

        c_aid = ws4.cell(row=row_idx, column=3, value=str(tx.account_id or ""))
        c_aid.font = font_normal
        c_aid.alignment = align_center

        c_anm = ws4.cell(row=row_idx, column=4, value=acc_map.get(tx.account_id, str(tx.account_id or "")))
        c_anm.font = font_normal
        c_anm.alignment = align_left

        c_typ = ws4.cell(row=row_idx, column=5, value=str(tx.type or ""))
        c_typ.font = font_normal
        c_typ.alignment = align_left

        c_cat = ws4.cell(row=row_idx, column=6, value=str(tx.category or "Altro"))
        c_cat.font = font_normal
        c_cat.alignment = align_left

        # Numeric: Importo (Right-Aligned)
        c_amt = ws4.cell(row=row_idx, column=7, value=_to_num(tx.amount))
        c_amt.number_format = FMT_CURRENCY
        c_amt.font = font_bold
        c_amt.alignment = align_right

        # Numeric: Commissioni (Right-Aligned)
        c_fee = ws4.cell(row=row_idx, column=8, value=_to_num(tx.fees, default=0.0))
        c_fee.number_format = FMT_CURRENCY_PLAIN
        c_fee.font = font_normal
        c_fee.alignment = align_right

        c_cur = ws4.cell(row=row_idx, column=9, value=str(tx.currency or "EUR"))
        c_cur.font = font_normal
        c_cur.alignment = align_center

        c_dsc = ws4.cell(row=row_idx, column=10, value=str(tx.description or ""))
        c_dsc.font = font_normal
        c_dsc.alignment = align_left

        c_sym = ws4.cell(row=row_idx, column=11, value=str(tx.asset_symbol or ""))
        c_sym.font = font_normal
        c_sym.alignment = align_center

        # Numeric: Quantità (Right-Aligned)
        if tx.quantity is not None:
            c_tx_qty = ws4.cell(row=row_idx, column=12, value=_to_num(tx.quantity))
            c_tx_qty.number_format = FMT_QTY
            c_tx_qty.font = font_normal
            c_tx_qty.alignment = align_right
        else:
            ws4.cell(row=row_idx, column=12, value="").font = font_normal

        # Numeric: Prezzo Unitario (Right-Aligned)
        if tx.unit_price is not None:
            c_tx_px = ws4.cell(row=row_idx, column=13, value=_to_num(tx.unit_price))
            c_tx_px.number_format = FMT_CURRENCY_PLAIN
            c_tx_px.font = font_normal
            c_tx_px.alignment = align_right
        else:
            ws4.cell(row=row_idx, column=13, value="").font = font_normal

        for c in range(1, 14):
            ws4.cell(row=row_idx, column=c).border = thin_border
            if row_idx % 2 == 0:
                ws4.cell(row=row_idx, column=c).fill = fill_alt
        current_row_tx += 1

    if transactions:
        tot_tx_row = current_row_tx
        c_tot_tx_lbl = ws4.cell(row=tot_tx_row, column=1, value="TOTALE FLUSSI & COMMISSIONI")
        c_tot_tx_lbl.font = font_total
        c_tot_tx_lbl.alignment = align_left

        for c_empty in range(2, 7):
            ws4.cell(row=tot_tx_row, column=c_empty, value="")

        t_amt = ws4.cell(row=tot_tx_row, column=7, value=f"=SUM(G{start_row_tx}:G{tot_tx_row - 1})")
        t_amt.number_format = FMT_CURRENCY
        t_amt.font = font_total
        t_amt.alignment = align_right

        t_fees = ws4.cell(row=tot_tx_row, column=8, value=f"=SUM(H{start_row_tx}:H{tot_tx_row - 1})")
        t_fees.number_format = FMT_CURRENCY_PLAIN
        t_fees.font = font_total
        t_fees.alignment = align_right

        for c_empty in range(9, 14):
            ws4.cell(row=tot_tx_row, column=c_empty, value="")

        for c in range(1, 14):
            ws4.cell(row=tot_tx_row, column=c).border = total_border
            ws4.cell(row=tot_tx_row, column=c).fill = fill_total

    # Auto-fit columns across all sheets safely (skipping titles in row 1 & 2)
    for ws in [ws1, ws2, ws3, ws4]:
        for col in ws.columns:
            max_len = 0
            for cell in col:
                if cell.row in [1, 2]:
                    continue
                val = cell.value
                if val is not None:
                    s = str(val)
                    if not s.startswith("="):
                        max_len = max(max_len, len(s))
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 13)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def generate_dataframe_excel(
    df: pd.DataFrame,
    title: str = "LUCAINVESTRATTO - Export Dati",
    sheet_name: str = "Dati",
    currency_cols: Optional[List[str]] = None,
    pct_cols: Optional[List[str]] = None,
    qty_cols: Optional[List[str]] = None,
    center_cols: Optional[List[str]] = None,
    sum_cols: Optional[List[str]] = None
) -> bytes:
    """
    Generates a beautifully formatted single-sheet Excel workbook (.xlsx) from a Pandas DataFrame.
    Guarantees:
    - ALL numeric data is native float/int (never text strings).
    - Explicit RIGHT ALIGNMENT on all numerical cells.
    - Standard professional Excel number formats (Currency, Percent, Quantity).
    - Institutional header, zebra stripes, and formula totals row.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.views.sheetView[0].showGridLines = True

    # Title & Subtitle
    ws["A1"] = title
    ws["A1"].font = FONT_TITLE
    ws["A2"] = f"Generato il: {datetime.now().strftime('%d/%m/%Y %H:%M')} | LUCAINVESTRATTO"
    ws["A2"].font = FONT_SUBTITLE

    c_cols = set(currency_cols or [])
    p_cols = set(pct_cols or [])
    q_cols = set(qty_cols or [])
    ctr_cols = set(center_cols or [])
    s_cols = set(sum_cols or [])

    # Headers at row 4
    for col_idx, col_name in enumerate(df.columns, 1):
        cell = ws.cell(row=4, column=col_idx, value=col_name)
        cell.font = FONT_HEADER
        cell.fill = FILL_HEADER
        cell.border = THIN_BORDER
        if col_name in c_cols or col_name in p_cols or col_name in q_cols or pd.api.types.is_numeric_dtype(df[col_name]):
            cell.alignment = ALIGN_HEADER_RIGHT
        elif col_name in ctr_cols:
            cell.alignment = ALIGN_HEADER_CENTER
        else:
            cell.alignment = ALIGN_HEADER_LEFT

    # Data rows from row 5
    for r_idx, row_data in enumerate(df.itertuples(index=False), 5):
        is_even = (r_idx - 4) % 2 == 0
        for c_idx, col_name in enumerate(df.columns, 1):
            raw_val = getattr(row_data, row_data._fields[c_idx - 1])
            cell = ws.cell(row=r_idx, column=c_idx)
            cell.border = THIN_BORDER
            if is_even:
                cell.fill = FILL_ALT

            if col_name in c_cols:
                cell.value = _to_num(raw_val)
                cell.number_format = FMT_CURRENCY
                cell.alignment = ALIGN_RIGHT
                cell.font = FONT_NORMAL
            elif col_name in p_cols:
                v = _to_num(raw_val)
                # Convert 15.5 (%) to 0.155 for Excel percentage display
                cell.value = v / 100.0 if abs(v) > 1.0 or v == 0.0 else v
                cell.number_format = FMT_PCT_SIGNED if "P&L" in col_name else FMT_PCT
                cell.alignment = ALIGN_RIGHT
                cell.font = FONT_NORMAL
            elif col_name in q_cols:
                cell.value = _to_num(raw_val)
                cell.number_format = FMT_QTY
                cell.alignment = ALIGN_RIGHT
                cell.font = FONT_NORMAL
            elif pd.api.types.is_numeric_dtype(df[col_name]):
                cell.value = _to_num(raw_val)
                cell.number_format = FMT_DECIMAL
                cell.alignment = ALIGN_RIGHT
                cell.font = FONT_NORMAL
            elif col_name in ctr_cols:
                cell.value = str(raw_val) if raw_val is not None and not pd.isna(raw_val) else ""
                cell.alignment = ALIGN_CENTER
                cell.font = FONT_NORMAL
            else:
                cell.value = str(raw_val) if raw_val is not None and not pd.isna(raw_val) else ""
                cell.alignment = ALIGN_LEFT
                cell.font = FONT_NORMAL

    # Total row
    if not df.empty and s_cols:
        start_row = 5
        end_row = 4 + len(df)
        tot_row = end_row + 1

        lbl_cell = ws.cell(row=tot_row, column=1, value="TOTALE")
        lbl_cell.font = FONT_TOTAL
        lbl_cell.alignment = ALIGN_LEFT
        lbl_cell.border = TOTAL_BORDER
        lbl_cell.fill = FILL_TOTAL

        for c_idx, col_name in enumerate(df.columns, 1):
            cell = ws.cell(row=tot_row, column=c_idx)
            cell.border = TOTAL_BORDER
            cell.fill = FILL_TOTAL
            col_letter = get_column_letter(c_idx)

            if col_name in s_cols:
                cell.value = f"=SUM({col_letter}{start_row}:{col_letter}{end_row})"
                if col_name in c_cols:
                    cell.number_format = FMT_CURRENCY
                elif col_name in q_cols:
                    cell.number_format = FMT_QTY
                else:
                    cell.number_format = FMT_DECIMAL
                cell.alignment = ALIGN_RIGHT
                cell.font = FONT_TOTAL
            elif c_idx > 1:
                if cell.value is None:
                    cell.value = ""

    # Auto-fit columns
    for col in ws.columns:
        max_len = 0
        for cell in col:
            if cell.row in [1, 2]:
                continue
            if cell.value is not None:
                s = str(cell.value)
                if not s.startswith("="):
                    max_len = max(max_len, len(s))
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 13)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
