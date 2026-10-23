"""
Unit tests to verify that all Excel (.xlsx) and CSV (.csv) exports
store numeric values strictly as numbers (float/int), NEVER as text strings.
"""
import unittest
import io
import openpyxl
import pandas as pd
import numpy as np

from lucainvestratto.models import BankAccount, Holding, Transaction
from lucainvestratto.export.excel_report import generate_excel_report
from lucainvestratto.export.csv_report import (
    generate_transactions_csv,
    generate_holdings_csv,
    generate_asset_allocation_csv,
    generate_risk_kpi_csv
)


class TestNumericExports(unittest.TestCase):
    def setUp(self):
        self.accounts = [
            BankAccount(id="acc_main", name="Conto Principale", initial_balance=5000.0),
            BankAccount(id="acc_invest", name="Dossier Titoli", initial_balance=15000.0)
        ]
        self.holdings = [
            Holding(
                id="h1",
                account_id="acc_invest",
                symbol="SWDA.MI",
                name="iShares Core MSCI World",
                asset_class="ETF / Fondi (ETFs & Funds)",
                sector="Azionario Globale",
                geography="Globale",
                currency="EUR",
                quantity=100.0,
                average_buy_price=80.0,
                current_price=105.0,
                annual_dividend_yield=1.5,
                duration=0.0
            ),
            Holding(
                id="h2",
                account_id="acc_invest",
                symbol="BTP37",
                name="BTP 4% 2037",
                asset_class="Obbligazionario (Bonds / Fixed Income)",
                sector="Governativo",
                geography="Italia",
                currency="EUR",
                quantity=50.0,
                average_buy_price=100.0,
                current_price=98.0,
                annual_dividend_yield=4.0,
                duration=9.5
            )
        ]
        self.transactions = [
            Transaction(
                id="tx1",
                account_id="acc_main",
                date="2026-01-15",
                type="Entrata (Inflow)",
                category="Stipendio / Reddito",
                amount=3200.0,
                currency="EUR",
                description="Accredito Stipendio",
                fees=0.0
            ),
            Transaction(
                id="tx2",
                account_id="acc_main",
                date="2026-01-20",
                type="Uscita (Outflow)",
                category="Spese Fisse / Utenze",
                amount=-850.50,
                currency="EUR",
                description="Affitto & Bollette",
                fees=1.50
            ),
            Transaction(
                id="tx3",
                account_id="acc_invest",
                date="2026-01-22",
                type="Acquisto Titoli (Buy)",
                category="Investimento Capitale",
                amount=-1050.0,
                currency="EUR",
                description="Acquisto SWDA",
                asset_symbol="SWDA.MI",
                quantity=10.0,
                unit_price=105.0,
                fees=2.95
            )
        ]

    # =========================================================================
    # EXCEL (.xlsx) NUMERIC INTEGRITY TESTS
    # =========================================================================
    def test_excel_export_cell_types(self):
        """Verify that every numeric cell in the multi-sheet Excel file has data_type == 'n' or 'f'."""
        xlsx_bytes = generate_excel_report(self.accounts, self.holdings, self.transactions)
        wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))

        expected_sheets = [
            "Sintesi & Metriche Rischio",
            "Asset Allocation",
            "Posizioni Portafoglio",
            "Estratto Conto Transazioni"
        ]
        for s in expected_sheets:
            self.assertIn(s, wb.sheetnames)

        # 1. Sheet 1: Sintesi & Metriche
        ws1 = wb["Sintesi & Metriche Rischio"]
        for row in range(5, 15):
            cell_b = ws1[f"B{row}"]
            if cell_b.value is not None:
                self.assertIn(cell_b.data_type, ['n', 'f'], f"Cell {cell_b.coordinate} must be numeric, got {cell_b.data_type}: {cell_b.value}")
                self.assertIsInstance(cell_b.value, (int, float), f"Cell {cell_b.coordinate} value must be int/float")
                self.assertEqual(cell_b.alignment.horizontal, 'right', f"Cell {cell_b.coordinate} must be right-aligned")
                self.assertFalse(str(cell_b.value).startswith("€"), f"Cell {cell_b.coordinate} contains raw currency string")

            cell_e = ws1[f"E{row}"]
            if cell_e.value is not None:
                self.assertIn(cell_e.data_type, ['n', 'f'], f"Cell {cell_e.coordinate} must be numeric, got {cell_e.data_type}: {cell_e.value}")
                self.assertIsInstance(cell_e.value, (int, float), f"Cell {cell_e.coordinate} value must be int/float")
                self.assertEqual(cell_e.alignment.horizontal, 'right', f"Cell {cell_e.coordinate} must be right-aligned")

        # 2. Sheet 2: Asset Allocation
        ws2 = wb["Asset Allocation"]
        for row in range(4, ws2.max_row + 1):
            for col in [2, 3, 4, 5, 6]:
                cell = ws2.cell(row=row, column=col)
                if cell.value is not None:
                    self.assertIn(cell.data_type, ['n', 'f'], f"Sheet2 {cell.coordinate} must be numeric/formula, got {cell.data_type}")
                    self.assertEqual(cell.alignment.horizontal, 'right', f"Sheet2 {cell.coordinate} must be right-aligned")
                    if cell.data_type == 'n':
                        self.assertIsInstance(cell.value, (int, float))

        # 3. Sheet 3: Posizioni Portafoglio
        ws3 = wb["Posizioni Portafoglio"]
        # Numeric columns: 7 (Quantità), 8 (PMC), 9 (Prezzo), 10 (Costo Totale), 11 (Controvalore),
        # 12 (P&L), 13 (Rendimento %), 14 (Peso %), 15 (Yield %), 16 (Reddito Annuo)
        numeric_cols_ws3 = [7, 8, 9, 10, 11, 12, 13, 14, 15, 16]
        for row in range(4, ws3.max_row + 1):
            for col in numeric_cols_ws3:
                cell = ws3.cell(row=row, column=col)
                if cell.value is not None and str(cell.value).strip() != "":
                    self.assertIn(cell.data_type, ['n', 'f'], f"Sheet3 {cell.coordinate} must be numeric/formula, got {cell.data_type}: {cell.value}")
                    self.assertEqual(cell.alignment.horizontal, 'right', f"Sheet3 {cell.coordinate} must be right-aligned")
                    if cell.data_type == 'n':
                        self.assertIsInstance(cell.value, (int, float))
                        self.assertFalse("€" in str(cell.value), f"Sheet3 {cell.coordinate} must not contain € text")

        # 4. Sheet 4: Estratto Conto Transazioni
        ws4 = wb["Estratto Conto Transazioni"]
        # Numeric columns: 7 (Importo), 8 (Commissioni), 12 (Quantità), 13 (Prezzo Unitario)
        numeric_cols_ws4 = [7, 8, 12, 13]
        for row in range(4, ws4.max_row + 1):
            for col in numeric_cols_ws4:
                cell = ws4.cell(row=row, column=col)
                if cell.value is not None and str(cell.value).strip() != "":
                    self.assertIn(cell.data_type, ['n', 'f'], f"Sheet4 {cell.coordinate} must be numeric/formula, got {cell.data_type}: {cell.value}")
                    self.assertEqual(cell.alignment.horizontal, 'right', f"Sheet4 {cell.coordinate} must be right-aligned")
                    if cell.data_type == 'n':
                        self.assertIsInstance(cell.value, (int, float))

    # =========================================================================
    # CSV (.csv) NUMERIC DTYPES & DATA INTEGRITY TESTS
    # =========================================================================
    def test_csv_transactions_numeric_types(self):
        """Verify that transactions CSV has purely numeric columns (float64)."""
        csv_bytes = generate_transactions_csv(self.transactions, self.accounts)
        df = pd.read_csv(io.BytesIO(csv_bytes))

        # Check column dtypes
        self.assertEqual(df["Importo"].dtype, np.float64)
        self.assertEqual(df["Commissioni"].dtype, np.float64)
        self.assertEqual(df["Quantita"].dtype, np.float64)
        self.assertEqual(df["Prezzo_Unitario"].dtype, np.float64)

        # Check values are exact floats
        self.assertAlmostEqual(df.loc[0, "Importo"], 3200.0)
        self.assertAlmostEqual(df.loc[1, "Importo"], -850.50)
        self.assertAlmostEqual(df.loc[1, "Commissioni"], 1.50)

        # Check that no string contains currency symbols
        for col in ["Importo", "Commissioni", "Quantita", "Prezzo_Unitario"]:
            for v in df[col].dropna():
                self.assertIsInstance(float(v), float)

    def test_csv_holdings_numeric_types(self):
        """Verify that holdings CSV has purely numeric columns (float64)."""
        csv_bytes = generate_holdings_csv(self.holdings)
        df = pd.read_csv(io.BytesIO(csv_bytes))

        numeric_cols = [
            "Quantita", "PMC_Prezzo_Medio_Carico", "Prezzo_Attuale",
            "Costo_Totale", "Controvalore_Attuale", "Plus_Minusvalenza_Latente",
            "Rendimento_Percentuale", "Peso_Portafoglio_Percentuale",
            "Yield_Dividendo_Percentuale", "Reddito_Annuo_Stimato", "Duration"
        ]
        for col in numeric_cols:
            self.assertEqual(df[col].dtype, np.float64, f"Column {col} must have float64 dtype")

        # Check calculated values
        swda_row = df[df["Simbolo_ISIN"] == "SWDA.MI"].iloc[0]
        self.assertAlmostEqual(swda_row["Quantita"], 100.0)
        self.assertAlmostEqual(swda_row["Costo_Totale"], 8000.0)
        self.assertAlmostEqual(swda_row["Controvalore_Attuale"], 10500.0)
        self.assertAlmostEqual(swda_row["Plus_Minusvalenza_Latente"], 2500.0)
        self.assertAlmostEqual(swda_row["Rendimento_Percentuale"], 31.25)

    def test_csv_asset_allocation_numeric_types(self):
        """Verify that asset allocation CSV has purely numeric columns (float64)."""
        csv_bytes = generate_asset_allocation_csv(self.holdings)
        df = pd.read_csv(io.BytesIO(csv_bytes))

        for col in ["Controvalore", "Costo_Totale", "Peso_Percentuale", "Plus_Minusvalenza", "Rendimento_Percentuale"]:
            self.assertEqual(df[col].dtype, np.float64, f"Column {col} must have float64 dtype")

    def test_csv_risk_kpi_numeric_types(self):
        """Verify that risk KPI CSV has strictly numeric values in Valore_Numerico."""
        csv_bytes = generate_risk_kpi_csv(self.holdings, self.accounts, self.transactions)
        df = pd.read_csv(io.BytesIO(csv_bytes))

        self.assertEqual(df["Valore_Numerico"].dtype, np.float64)
        self.assertGreater(len(df), 15)

    def test_csv_italian_excel_delimiter(self):
        """Verify that Italian Excel CSV (; delimiter and , decimal) is correctly parsed."""
        csv_bytes = generate_transactions_csv(self.transactions, self.accounts, delimiter=";", decimal=",")
        df = pd.read_csv(io.BytesIO(csv_bytes), sep=";", decimal=",")
        self.assertEqual(df["Importo"].dtype, np.float64)
        self.assertAlmostEqual(df.loc[1, "Importo"], -850.50)


if __name__ == "__main__":
    unittest.main()
